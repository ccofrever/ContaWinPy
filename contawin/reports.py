"""Informes contables (lógica pura, sin interfaz).

Cada función devuelve un objeto `Informe` que luego se puede:
  - imprimir / ver en vista previa / guardar en PDF  (ui/impresion.py, con Qt)
  - exportar a Excel (.xlsx con openpyxl) o CSV        (exportar_excel)

Equivalencias con el original:
    Lib_Diario()   LIBRO.PRG     -> libro_diario()
    Lib_DxTipo()   LIBROXT.PRG   -> libro_diario(tipo=...)
    Lib_Mayor()    MAYOR.PRG     -> libro_mayor()
    Bal8Win()      BAL8WIN.PRG   -> balance_8_columnas()
    Bal_TInfor()   BALTI.PRG     -> balance_tipo_informe()
    LibroC()       LIBROC.PRG    -> libro_compras()
    Lo_Imprime()   ASIENTOS.PRG  -> comprobante()
    Imprime*()     M*.PRG        -> listado_*()
"""
from __future__ import annotations

import csv
import datetime as _dt
from collections import defaultdict
from dataclasses import dataclass, field

from . import util
from .db import Database

# Estilos de fila
NORMAL, GRUPO, SUBTOTAL, TOTAL, ENCABEZADO_ASIENTO = "", "grupo", "subtotal", "total", "asiento"


@dataclass
class Columna:
    titulo: str
    ancho: float = 1.0          # peso relativo
    formato: str = "t"          # t texto, m monto, f fecha, c código cuenta, r rut
    alinear: str = ""           # L/R/C (por defecto: montos a la derecha)
    vacio_si_cero: bool = True

    @property
    def alineacion(self) -> str:
        return self.alinear or ("R" if self.formato == "m" else "L")


@dataclass
class Fila:
    valores: list
    estilo: str = NORMAL
    texto: str = ""             # para filas de grupo: texto que ocupa toda la línea


@dataclass
class Informe:
    titulo: str
    columnas: list[Columna]
    filas: list[Fila] = field(default_factory=list)
    subtitulos: list[str] = field(default_factory=list)
    membrete: list[str] = field(default_factory=list)        # empresa, rut, ciudad (arriba a la izquierda)
    datos_cabecera: list[str] = field(default_factory=list)   # líneas extra bajo el título
    pie: list[str] = field(default_factory=list)
    firmas: list[str] = field(default_factory=list)           # textos bajo líneas de firma
    horizontal: bool = False
    fecha_emision: _dt.date | None = None
    totales_pagina: list[int] = field(default_factory=list)   # columnas a totalizar por página
    nombre_archivo: str = "informe"

    def texto_celda(self, i: int, valor) -> str:
        c = self.columnas[i]
        if valor is None or valor == "":
            return ""
        if c.formato == "m":
            return util.fmt_monto(valor, c.vacio_si_cero)
        if c.formato == "f":
            return util.fmt_fecha(valor)
        if c.formato == "c":
            return util.formato_codigo(valor)
        if c.formato == "r":
            return util.formato_rut(valor)
        return str(valor)

    def textos(self, fila: Fila) -> list[str]:
        return [self.texto_celda(i, v) for i, v in enumerate(fila.valores)]


def _membrete(emp) -> list[str]:
    if not emp:
        return []
    return [emp["razon_social"], "R.U.T. " + util.formato_rut(emp["rut"]), emp["ciudad"]]


def _rango(desde, hasta) -> str:
    return f"DESDE EL < {util.fmt_fecha(desde)} > HASTA EL < {util.fmt_fecha(hasta)} >"


# ---------------------------------------------------------------------------
# Comprobante contable
# ---------------------------------------------------------------------------
def comprobante(db: Database, empresa_id: int, asiento_id: int, emision=None) -> Informe:
    emp = db.empresa(empresa_id)
    a = db.asiento(asiento_id)
    lineas = db.detalle_asiento(asiento_id)
    inf = Informe(
        titulo=f"COMPROBANTE  {util.nombre_tipo_asiento(a['tipo']).upper()}",
        membrete=_membrete(emp),
        datos_cabecera=[
            f"Folio N° : {util.fmt_monto(a['numero'])}        Fecha : {util.fmt_fecha(a['fecha'])}",
            f"GLOSA : {a['glosa']}",
            f"CENTRO DE COSTO : {db.nombre_ccosto(empresa_id, a['cdcosto']) or a['cdcosto']}",
        ],
        columnas=[Columna("N°", 0.5, alinear="C"), Columna("CÓDIGO", 1.1, "c"),
                  Columna("DESCRIPCIÓN CUENTA", 4), Columna("DEBE", 1.6, "m"), Columna("HABER", 1.6, "m")],
        fecha_emision=emision or _dt.date.today(),
        nombre_archivo=f"comprobante_{a['numero']}",
        firmas=["DIGITADO", "V°B° CONTABILIDAD"],
    )
    td = th = 0
    detalles_doc = []
    for i, l in enumerate(lineas, 1):
        inf.filas.append(Fila([f"{i:02d}", l["codigo"], db.nombre_cuenta(empresa_id, l["codigo"]),
                               l["debe"], l["haber"]]))
        td += l["debe"]
        th += l["haber"]
        d = l["documento"]
        if d:
            detalles_doc.append(
                f"{util.sigla_documento(d['tdocum'])} N° {util.fmt_monto(d['numero_doc'])} del "
                f"{util.fmt_fecha(d['fecha_doc'])} - {util.formato_rut(d['rut_prov'])} "
                f"{db.nombre_proveedor(empresa_id, d['rut_prov'])}"
                + (f" - {d['detalle']}" if d.get("detalle") else ""))
    inf.filas.append(Fila(["", "", "T O T A L   A S I E N T O", td, th], TOTAL))
    if detalles_doc:
        inf.pie = ["Documentos:"] + ["   " + t for t in detalles_doc]
    return inf


# ---------------------------------------------------------------------------
# Libro diario (y libro diario por tipo)
# ---------------------------------------------------------------------------
def libro_diario(db: Database, empresa_id: int, periodo_id: int, desde, hasta,
                 tipo: str | None = None, emision=None) -> Informe:
    emp = db.empresa(empresa_id)
    nombres = {c["codigo"]: c["nombre"] for c in db.cuentas(empresa_id)}
    sql = ["SELECT * FROM asiento WHERE periodo_id=? AND fecha>=? AND fecha<=?"]
    p = [periodo_id, util.to_iso(desde), util.to_iso(hasta)]
    if tipo:
        sql.append("AND tipo=?")
        p.append(tipo)
        sql.append("ORDER BY tipo, fecha, numero")
    else:
        sql.append("ORDER BY fecha, numero")
    asientos = db.q(" ".join(sql), p)

    titulo = "LIBRO DIARIO" + (f" POR TIPO  ({util.nombre_tipo_asiento(tipo).upper()})" if tipo else "")
    inf = Informe(
        titulo=titulo, subtitulos=[_rango(desde, hasta)], membrete=_membrete(emp),
        columnas=[Columna("Número", 0.8, alinear="R"), Columna("Fecha", 1, "f"), Columna("Cuenta", 0.9, "c"),
                  Columna("Nombre / Glosa", 3.6), Columna("D E B E", 1.3, "m"), Columna("H A B E R", 1.3, "m"),
                  Columna("Tot. DEBE", 1.3, "m"), Columna("Tot. HABER", 1.3, "m")],
        fecha_emision=emision or _dt.date.today(), totales_pagina=[6, 7],
        nombre_archivo="libro_diario" + (f"_{tipo}" if tipo else ""),
    )
    gd = gh = 0
    for a in asientos:
        inf.filas.append(Fila([f"{a['numero']:06d}", a["fecha"], "", a["glosa"], "", "", a["debe"], a["haber"]],
                              ENCABEZADO_ASIENTO))
        for d in db.q("SELECT * FROM detalle WHERE asiento_id=? ORDER BY linea, id", (a["id"],)):
            inf.filas.append(Fila(["", "", d["codigo"], nombres.get(d["codigo"], ""), d["debe"], d["haber"],
                                   "", ""]))
        gd += a["debe"]
        gh += a["haber"]
    inf.filas.append(Fila(["", "", "", "TOTAL ..............", "", "", gd, gh], TOTAL))
    return inf


# ---------------------------------------------------------------------------
# Libro mayor
# ---------------------------------------------------------------------------
def _saldo_txt(saldo: int) -> str:
    if saldo == 0:
        return "0"
    return f"{util.fmt_monto(abs(saldo))} {'D' if saldo > 0 else 'A'}"


def libro_mayor(db: Database, empresa_id: int, periodo_id: int, desde, hasta,
                cod_desde: str | None = None, cod_hasta: str | None = None, emision=None) -> Informe:
    """Movimientos por cuenta con saldo anterior (arrastre) y saldo acumulado.

    El saldo se muestra como Deudor (D) o Acreedor (A) = Debe - Haber.
    """
    emp = db.empresa(empresa_id)
    desde_iso, hasta_iso = util.to_iso(desde), util.to_iso(hasta)
    cuentas = db.cuentas(empresa_id)
    if cod_desde:
        cuentas = [c for c in cuentas if c["codigo"] >= cod_desde]
    if cod_hasta:
        cuentas = [c for c in cuentas if c["codigo"] <= cod_hasta]
    inf = Informe(
        titulo="LIBRO MAYOR", subtitulos=[_rango(desde, hasta)], membrete=_membrete(emp),
        columnas=[Columna("Número", 0.8, alinear="R"), Columna("Fecha", 1, "f"), Columna("G l o s a", 4.2),
                  Columna("D e b e", 1.4, "m"), Columna("H a b e r", 1.4, "m"),
                  Columna("S a l d o", 1.6, alinear="R")],
        fecha_emision=emision or _dt.date.today(), nombre_archivo="libro_mayor",
    )
    for c in cuentas:
        movs = db.q("""SELECT a.numero, a.fecha, a.glosa, d.debe, d.haber FROM detalle d
                       JOIN asiento a ON a.id=d.asiento_id
                       WHERE a.periodo_id=? AND d.codigo=? AND a.fecha>=? AND a.fecha<=?
                       ORDER BY a.fecha, a.numero, d.linea""", (periodo_id, c["codigo"], desde_iso, hasta_iso))
        if not movs:
            continue
        ant = db.q1("""SELECT COALESCE(SUM(d.debe),0) d, COALESCE(SUM(d.haber),0) h FROM detalle d
                       JOIN asiento a ON a.id=d.asiento_id
                       WHERE a.periodo_id=? AND d.codigo=? AND a.fecha<?""", (periodo_id, c["codigo"], desde_iso))
        saldo = ant["d"] - ant["h"]
        inf.filas.append(Fila([], GRUPO, f"{util.formato_codigo(c['codigo'])}   {c['nombre']}"))
        inf.filas.append(Fila(["", "", " A R R A S T R E  (saldo anterior)", ant["d"], ant["h"],
                               _saldo_txt(saldo)], SUBTOTAL))
        td = th = 0
        for m in movs:
            saldo += m["debe"] - m["haber"]
            td += m["debe"]
            th += m["haber"]
            inf.filas.append(Fila([util.fmt_monto(m["numero"]), m["fecha"], m["glosa"], m["debe"], m["haber"],
                                   _saldo_txt(saldo)]))
        inf.filas.append(Fila(["", "", " TOTAL PERÍODO / S A L D O", td, th, _saldo_txt(saldo)], TOTAL))
    return inf


# ---------------------------------------------------------------------------
# Balance de 8 columnas
# ---------------------------------------------------------------------------
def _columna_cuenta(tributario: bool, ancho: float, titulo_codigo: str) -> Columna:
    """Balance tributario (para bancos u otros terceros): un correlativo en lugar del código de
    cuenta, para no exponer el plan de cuentas. Borrador: el código, para el trabajo interno."""
    if tributario:
        return Columna("N°", ancho, "t", "R")
    return Columna(titulo_codigo, ancho, "c")


TEXTO_ART100 = [
    "Artículo 100 Código Tributario: Dejo constancia que la contabilidad de este ejercicio, así como el "
    "Inventario y Balance",
    "correspondiente, han sido confeccionados por el contador en base a datos fidedignos que le han "
    "proporcionado.",
]


def calcular_balance_8(db: Database, empresa_id: int, periodo_id: int, desde, hasta) -> dict:
    filas = []
    tot = defaultdict(int)
    sums = {r["codigo"]: (r["d"], r["h"]) for r in db.q(
        """SELECT d.codigo, SUM(d.debe) d, SUM(d.haber) h FROM detalle d JOIN asiento a ON a.id=d.asiento_id
           WHERE a.periodo_id=? AND a.fecha>=? AND a.fecha<=? GROUP BY d.codigo""",
        (periodo_id, util.to_iso(desde), util.to_iso(hasta)))}
    nombres = {c["codigo"]: c["nombre"] for c in db.cuentas(empresa_id)}
    for codigo in sorted(set(sums)):
        deb, cre = sums[codigo]
        dif = deb - cre
        deudor, acreedor = (dif, 0) if dif > 0 else (0, -dif)
        activo = pasivo = perdida = ganancia = 0
        if util.es_cuenta_resultado(codigo):
            perdida, ganancia = deudor, acreedor
        else:
            activo, pasivo = deudor, acreedor
        f = dict(codigo=codigo, nombre=nombres.get(codigo, "(cuenta no existe)"), debitos=deb, creditos=cre,
                 deudor=deudor, acreedor=acreedor, activo=activo, pasivo=pasivo, perdida=perdida,
                 ganancia=ganancia)
        filas.append(f)
        for k in ("debitos", "creditos", "deudor", "acreedor", "activo", "pasivo", "perdida", "ganancia"):
            tot[k] += f[k]
    # Resultado del ejercicio
    cal2 = tot["ganancia"] - tot["perdida"]
    cal3 = tot["activo"] - tot["pasivo"]
    res = dict(activo=0, pasivo=0, perdida=0, ganancia=0)
    if cal2 > 0:
        res["perdida"] = cal2
    else:
        res["ganancia"] = -cal2
    if cal3 > 0:
        res["pasivo"] = cal3
    else:
        res["activo"] = -cal3
    es_ganancia = cal3 > 0
    sumas = dict(tot)
    for k in ("activo", "pasivo", "perdida", "ganancia"):
        sumas[k] = tot[k] + res[k]
    return dict(filas=filas, totales=dict(tot), resultado=res, es_ganancia=es_ganancia, sumas=sumas)


def balance_8_columnas(db: Database, empresa_id: int, periodo_id: int, desde, hasta, emision=None,
                       tributario: bool = False) -> Informe:
    emp = db.empresa(empresa_id)
    b = calcular_balance_8(db, empresa_id, periodo_id, desde, hasta)
    claves = ["debitos", "creditos", "deudor", "acreedor", "activo", "pasivo", "perdida", "ganancia"]
    inf = Informe(
        titulo="B A L A N C E    G E N E R A L",
        subtitulos=[f"Ejercicio  DESDE : {util.fmt_fecha(desde)}   HASTA : {util.fmt_fecha(hasta)}"],
        datos_cabecera=[f"Empresa   : {emp['razon_social']}", f"R.U.T.    : {util.formato_rut(emp['rut'])}",
                        f"Dirección : {emp['direccion']}", f"Ciudad    : {emp['ciudad']}",
                        f"Giro      : {emp['giro']}",
                        "S A L D O S  (Deudor / Acreedor)  ·  I N V E N T A R I O  (Activo / Pasivo)  ·  "
                        "R E S U L T A D O  (Pérdida / Ganancia)"],
        columnas=[_columna_cuenta(tributario, 0.9, "CÓDIGO"), Columna("C U E N T A", 2.4)] +
                 [Columna(t, 1.15, "m") for t in ("DÉBITOS", "CRÉDITOS", "DEUDOR", "ACREEDOR", "ACTIVO",
                                                    "PASIVO", "PÉRDIDA", "GANANCIA")],
        horizontal=True, fecha_emision=emision or _dt.date.today(),
        nombre_archivo="balance_8_columnas_tributario" if tributario else "balance_8_columnas",
        pie=TEXTO_ART100, firmas=["CONTADOR", "CONTRIBUYENTE O REPRESENTANTE LEGAL"],
    )
    for n, f in enumerate(b["filas"], 1):
        inf.filas.append(Fila([n if tributario else f["codigo"], f["nombre"]] + [f[k] for k in claves]))
    t = b["totales"]
    inf.filas.append(Fila(["", "T O T A L E S"] + [t.get(k, 0) for k in claves], TOTAL))
    r = b["resultado"]
    inf.filas.append(Fila(["", "G A N A N C I A" if b["es_ganancia"] else "P E R D I D A", "", "", "", "",
                           r["activo"], r["pasivo"], r["perdida"], r["ganancia"]], SUBTOTAL))
    s = b["sumas"]
    inf.filas.append(Fila(["", "S U M A S"] + [s.get(k, 0) for k in claves], TOTAL))
    return inf


# ---------------------------------------------------------------------------
# Balance tipo informe
# ---------------------------------------------------------------------------
def balance_tipo_informe(db: Database, empresa_id: int, periodo_id: int, hasta, emision=None,
                         tributario: bool = False) -> Informe:
    """Saldo de cada cuenta desde el inicio del año hasta la fecha, agrupado por
    el primer dígito del código. Grupo 1: Debe - Haber; resto: Haber - Debe."""
    emp = db.empresa(empresa_id)
    sums = {r["codigo"]: (r["d"], r["h"]) for r in db.q(
        """SELECT d.codigo, SUM(d.debe) d, SUM(d.haber) h FROM detalle d JOIN asiento a ON a.id=d.asiento_id
           WHERE a.periodo_id=? AND a.fecha<=? GROUP BY d.codigo""", (periodo_id, util.to_iso(hasta)))}
    inf = Informe(
        titulo=f"BALANCE TIPO INFORME   HASTA EL < {util.fmt_fecha(hasta)} >", subtitulos=[emp["razon_social"]],
        membrete=_membrete(emp),
        columnas=[_columna_cuenta(tributario, 1, "Código"), Columna("Cuenta", 4), Columna("S a l d o", 1.6, "m"),
                  Columna("SubTotal", 1.6, "m")],
        fecha_emision=emision or _dt.date.today(), nombre_archivo="balance_tipo_informe_tributario" if tributario else "balance_tipo_informe",
    )
    grupo_actual, sub, sub_tiene, n = None, 0, False, 0
    for c in db.cuentas(empresa_id):
        d, h = sums.get(c["codigo"], (0, 0))
        g = c["codigo"][:1]
        if grupo_actual is not None and g != grupo_actual and sub_tiene:
            inf.filas.append(Fila(["", "", "", sub], SUBTOTAL))
        if g != grupo_actual:
            grupo_actual, sub, sub_tiene = g, 0, False
        total = (d - h) if g == "1" else (h - d)
        if total != 0:
            n += 1
            inf.filas.append(Fila([n if tributario else c["codigo"], c["nombre"], total, ""]))
            sub += total
            sub_tiene = True
    if grupo_actual is not None and sub_tiene:
        inf.filas.append(Fila(["", "", "", sub], SUBTOTAL))
    return inf


# ---------------------------------------------------------------------------
# Libro de compras
# ---------------------------------------------------------------------------
def neto_iva_total(c) -> tuple[int, int, int]:
    neto, iva, total = c["neto"], c["iva"], c["total"]
    if not neto and not iva:
        neto = total - (c["adicional"] if "adicional" in c.keys() else 0)
    return neto, iva, total


def libro_compras(db: Database, empresa_id: int, periodo_id: int, desde, hasta,
                  cdcosto: str | None = None, emision=None) -> Informe:
    """cdcosto=None -> todos los centros de costo (cada uno con su subtotal)."""
    emp = db.empresa(empresa_id)
    compras = db.compras(periodo_id, desde, hasta, cdcosto)
    proveedores = {p["rut"]: p["nombre"] for p in db.proveedores(empresa_id)}
    ccostos = {c["codigo"]: c["nombre"] for c in db.ccostos(empresa_id)}
    inf = Informe(
        titulo="LIBRO DE COMPRAS", subtitulos=[_rango(desde, hasta)], membrete=_membrete(emp),
        columnas=[Columna("Tipo Doc.", 0.8), Columna("Cuenta", 0.8, "c"), Columna("N° Documento", 1, "m"),
                  Columna("Fecha", 0.8, "f"), Columna("R.U.T.", 1, "r"), Columna("Nombre", 2.4),
                  Columna("Detalle", 2.4), Columna("NETO", 1, "m", vacio_si_cero=False),
                  Columna("I.V.A.", 0.9, "m", vacio_si_cero=False), Columna("TOTAL", 1, "m", vacio_si_cero=False),
                  Columna("Fec. Pago", 0.8, "f")],
        horizontal=True, fecha_emision=emision or _dt.date.today(), nombre_archivo="libro_compras",
    )
    grupos: dict[str, list] = defaultdict(list)
    for c in compras:
        grupos[c["cdcosto"] or ""].append(c)
    gn = gi = gt = 0
    for cc in sorted(grupos):
        nom = ccostos.get(cc, "") if cc else "SIN CENTRO DE COSTO"
        inf.filas.append(Fila([], GRUPO, f"Centro de Costos  {cc}   {nom}"))
        pn = pi = pt = 0
        for c in grupos[cc]:
            neto, iva, total = neto_iva_total(c)
            inf.filas.append(Fila([util.sigla_documento(c["tdocum"]), c["ccuenta"], c["numero_doc"],
                                   c["fecha_doc"], c["rut_prov"], proveedores.get(c["rut_prov"], ""),
                                   c["detalle"], neto, iva, total, c["fecha_pago"]]))
            pn, pi, pt = pn + neto, pi + iva, pt + total
        inf.filas.append(Fila(["", "", "", "", "", "", "TOTAL CENTRO DE COSTO -->", pn, pi, pt, ""], SUBTOTAL))
        gn, gi, gt = gn + pn, gi + pi, gt + pt
    inf.filas.append(Fila(["", "", "", "", "", "", "TOTAL GENERAL -->", gn, gi, gt, ""], TOTAL))
    return inf


# ---------------------------------------------------------------------------
# Listados de mantenedores
# ---------------------------------------------------------------------------
def listado_empresas(db: Database) -> Informe:
    inf = Informe(titulo="INFORME GENERAL DATOS EMPRESAS",
                  columnas=[Columna("Razón Social", 3), Columna("R.U.T.", 1.2, "r"), Columna("Dirección", 3),
                            Columna("Ciudad", 1.4), Columna("Directorio", 1)],
                  fecha_emision=_dt.date.today(), nombre_archivo="empresas")
    for e in db.empresas():
        inf.filas.append(Fila([e["razon_social"], e["rut"], e["direccion"], e["ciudad"], e["directorio"]]))
    return inf


def listado_cuentas(db: Database, empresa_id: int) -> Informe:
    emp = db.empresa(empresa_id)
    inf = Informe(titulo="INFORME GENERAL PLAN DE CUENTAS", membrete=_membrete(emp),
                  columnas=[Columna("CÓDIGO", 1, "c"), Columna("Nombre", 4),
                            Columna("Pide documento", 1.2, alinear="C")],
                  fecha_emision=_dt.date.today(), nombre_archivo="plan_de_cuentas")
    for c in db.cuentas(empresa_id):
        inf.filas.append(Fila([c["codigo"], c["nombre"], "Sí" if c["cdocum"] else ""]))
    return inf


def listado_ccostos(db: Database, empresa_id: int) -> Informe:
    emp = db.empresa(empresa_id)
    inf = Informe(titulo="INFORME GENERAL CENTRO DE COSTO", membrete=_membrete(emp),
                  columnas=[Columna("CÓDIGO", 1), Columna("Nombre", 4)],
                  fecha_emision=_dt.date.today(), nombre_archivo="centros_de_costo")
    for c in db.ccostos(empresa_id):
        inf.filas.append(Fila([c["codigo"], c["nombre"]]))
    return inf


def listado_proveedores(db: Database, empresa_id: int) -> Informe:
    emp = db.empresa(empresa_id)
    inf = Informe(titulo="INFORME GENERAL PROVEEDORES", membrete=_membrete(emp),
                  columnas=[Columna("R.U.T.", 1.2, "r"), Columna("Nombre", 3), Columna("Dirección", 2.4),
                            Columna("Ciudad", 1.3), Columna("Teléfono", 1.2), Columna("E-mail", 1.8)],
                  horizontal=True, fecha_emision=_dt.date.today(), nombre_archivo="proveedores")
    for p in db.proveedores(empresa_id, "nombre"):
        inf.filas.append(Fila([p["rut"], p["nombre"], p["direccion"], p["ciudad"], p["telefono"], p["email"]]))
    return inf


# ---------------------------------------------------------------------------
# Exportación a Excel / CSV
# ---------------------------------------------------------------------------
def exportar_excel(inf: Informe, ruta: str) -> str:
    """Exporta a .xlsx (si openpyxl está instalado) o .csv. Devuelve la ruta final."""
    if ruta.lower().endswith(".xlsx"):
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            ruta = ruta[:-5] + ".csv"
    if ruta.lower().endswith(".csv"):
        with open(ruta, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f, delimiter=";")
            w.writerow([inf.titulo])
            for s in inf.membrete + inf.subtitulos:
                w.writerow([s])
            w.writerow([])
            w.writerow([c.titulo for c in inf.columnas])
            for fila in inf.filas:
                if fila.estilo == GRUPO:
                    w.writerow([fila.texto])
                    continue
                w.writerow([_valor_excel(inf, i, v) for i, v in enumerate(fila.valores)])
        return ruta

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = inf.nombre_archivo[:30] or "Informe"
    ws.append([inf.titulo])
    ws["A1"].font = Font(name="Arial", size=14, bold=True)
    for s in inf.membrete + inf.subtitulos:
        ws.append([s])
    ws.append([])
    ws.append([c.titulo for c in inf.columnas])
    fila_titulos = ws.max_row
    gris = PatternFill("solid", fgColor="D9D9D9")
    for i in range(1, len(inf.columnas) + 1):
        cel = ws.cell(row=fila_titulos, column=i)
        cel.font = Font(bold=True)
        cel.fill = gris
        cel.border = Border(bottom=Side(style="thin"))
    for fila in inf.filas:
        if fila.estilo == GRUPO:
            ws.append([fila.texto])
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True, color="1F3864")
            continue
        ws.append([_valor_excel(inf, i, v) for i, v in enumerate(fila.valores)])
        r = ws.max_row
        for i, c in enumerate(inf.columnas, 1):
            cel = ws.cell(row=r, column=i)
            if c.formato == "m" and isinstance(cel.value, (int, float)):
                cel.number_format = "#,##0"
            elif c.formato == "f" and isinstance(cel.value, _dt.date):
                cel.number_format = "DD/MM/YYYY"
            if c.alineacion == "R":
                cel.alignment = Alignment(horizontal="right")
            if fila.estilo in (TOTAL, SUBTOTAL, ENCABEZADO_ASIENTO):
                cel.font = Font(bold=True)
    for i, c in enumerate(inf.columnas, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(10, min(60, int(c.ancho * 12)))
    ws.freeze_panes = ws.cell(row=fila_titulos + 1, column=1)
    wb.save(ruta)
    return ruta


def _valor_excel(inf: Informe, i: int, v):
    c = inf.columnas[i]
    if v is None or v == "":
        return None
    if c.formato == "m":
        if isinstance(v, (int, float)):
            return int(v)
        return v
    if c.formato == "f":
        return util.from_iso(v) or v
    return inf.texto_celda(i, v)
